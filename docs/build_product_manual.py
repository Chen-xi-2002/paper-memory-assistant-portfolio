#!/usr/bin/env python3
"""Build the formatted product manual for the paper memory assistant."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "论文记忆与引用助手产品说明书.docx"

INK = "22302C"
MUTED = "66736E"
FOREST = "175B4D"
FOREST_DARK = "0F443A"
OCHRE = "B7793F"
PAPER = "FBFAF6"
LINE = "DCE3DE"
SOFT_GREEN = "EAF3EF"
SOFT_OCHRE = "F8EEE1"
WHITE = "FFFFFF"
BODY_FONT = "Hiragino Sans GB"
HEAD_FONT = "Hiragino Sans GB"


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def set_style_font(style, font_name: str, size: float, bold: bool | None = None, color: str | None = None) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def remove_paragraph_borders(paragraph_or_style) -> None:
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_east_asia_font(run, BODY_FONT)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_element = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run_element.append(text)
    field.append(run_element)
    paragraph._p.append(field)

    tail = paragraph.add_run(" 页")
    set_east_asia_font(tail, BODY_FONT)
    tail.font.size = Pt(9)
    tail.font.color.rgb = RGBColor.from_string(MUTED)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), FOREST)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        set_east_asia_font(lead, BODY_FONT)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_east_asia_font(rest, BODY_FONT)
    else:
        run = paragraph.add_run(text)
        set_east_asia_font(run, BODY_FONT)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.28


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_east_asia_font(run, BODY_FONT)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.2


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, SOFT_GREEN)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(FOREST_DARK)
        set_east_asia_font(run, BODY_FONT)
        if widths:
            cell.width = Inches(widths[index])
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cell = cells[index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(value)
            run.font.size = Pt(9.2)
            run.font.color.rgb = RGBColor.from_string(INK)
            set_east_asia_font(run, BODY_FONT)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_section_title(doc: Document, title: str) -> None:
    paragraph = doc.add_heading(title, level=1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(7)


def add_subheading(doc: Document, title: str) -> None:
    paragraph = doc.add_heading(title, level=2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(11)
    paragraph.paragraph_format.space_after = Pt(5)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    set_style_font(styles["Normal"], BODY_FONT, 10.8, False, INK)
    styles["Normal"].paragraph_format.line_spacing = 1.28
    styles["Normal"].paragraph_format.space_after = Pt(6)

    set_style_font(styles["Title"], HEAD_FONT, 28, True, FOREST_DARK)
    styles["Title"].paragraph_format.space_after = Pt(8)
    remove_paragraph_borders(styles["Title"])

    set_style_font(styles["Heading 1"], HEAD_FONT, 17, True, FOREST_DARK)
    set_style_font(styles["Heading 2"], HEAD_FONT, 13, True, INK)
    set_style_font(styles["Heading 3"], HEAD_FONT, 11.3, True, FOREST)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    title = doc.add_paragraph("论文记忆与引用助手产品说明书", style="Title")
    title.paragraph_format.space_before = Pt(54)
    title.paragraph_format.space_after = Pt(9)
    remove_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("面向论文阅读、记忆、全文查证与可追溯引用的本地优先 Web 应用")
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)
    set_east_asia_font(subtitle_run, BODY_FONT)
    subtitle.paragraph_format.space_after = Pt(28)

    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cover_table.autofit = False
    cover_rows = [
        ("文档版本", "v1.0 个人项目初稿"),
        ("产品阶段", "可交互原型 / 求职项目"),
        ("适用对象", "论文写作者、产品经理面试官、项目协作者"),
        ("文档用途", "说明产品定位、功能边界、使用方法和验收标准"),
    ]
    for row, (label, value) in zip(cover_table.rows, cover_rows):
        for cell in row.cells:
            set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
        row.cells[0].width = Inches(1.15)
        row.cells[1].width = Inches(5.65)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(9.2)
        label_run.font.color.rgb = RGBColor.from_string(FOREST)
        set_east_asia_font(label_run, BODY_FONT)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        value_run.font.size = Pt(9.2)
        value_run.font.color.rgb = RGBColor.from_string(INK)
        set_east_asia_font(value_run, BODY_FONT)
    set_table_borders(cover_table, color="E7EBE8", size=3)
    set_repeat_table_header(cover_table.rows[0])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_section_title(doc, "1 文档目的")
    add_paragraph(doc, "本说明书用于帮助使用者理解论文记忆与引用助手的定位、核心流程、功能边界和使用方法。它也是项目交付与面试演示的配套材料，重点说明产品如何把论文、笔记和关键证据转化为可检索、可核对的写作资产。")
    add_paragraph(doc, "本说明书面向正在撰写论文的学生、需要评估产品能力的招聘方，以及后续参与迭代的产品、设计和研发协作者。阅读完成后，使用者应能独立完成论文导入、问题检索、单篇全文查证和在线候选论文处理。")
    add_paragraph(doc, "当前版本为个人项目原型。它已经具备核心闭环，但仍需要真实用户访谈、可用性测试和检索基准数据来验证产品假设。文档中的目标指标用于验证计划，不代表已经取得的结果。")

    add_section_title(doc, "2 产品概述")
    add_subheading(doc, "2.1 产品定位")
    add_paragraph(doc, "论文记忆与引用助手是一个本地优先的证据记忆与检索工具。它不试图替代论文阅读，也不生成无法核对的引用；它解决的是读过的论文、笔记和证据在写作时无法快速找回的问题。")
    add_table(
        doc,
        ["维度", "定义"],
        [
            ["目标问题", "论文记忆断层、笔记分散、引用核验耗时、通用问答缺少来源"],
            ["目标用户", "同时阅读多篇论文并需要频繁查找论据的研究生和学术写作者"],
            ["核心能力", "导入论文、保存笔记、全文检索、证据问答、单篇查证、在线发现"],
            ["差异点", "回答区分全文、笔记、标注、摘要和待补 PDF，证据不足时明确说明"],
            ["数据策略", "论文和笔记默认保存在用户浏览器本地，可导出 JSON 备份"],
        ],
        widths=[1.25, 5.55],
    )

    add_subheading(doc, "2.2 用户问题")
    add_bullets(doc, [
        "读过很多论文，但写作时只记得观点，不记得出自哪篇论文。",
        "笔记散落在 PDF 批注、文档和聊天工具中，无法按写作问题统一检索。",
        "通用问答工具可能给出流畅答案，却无法稳定回到真正读过的原文。",
        "在线找到的候选论文缺少全文时，用户无法判断摘要能否支持具体结论。",
        "现有文献管理工具擅长存储和格式化引用，但对个人笔记与全文问答支持有限。",
    ])

    add_subheading(doc, "2.3 产品价值")
    add_paragraph(doc, "产品价值不在于一次回答多少问题，而在于让每一条回答都能回到可核对的材料。使用者可以先确认来源，再决定是否引用；当材料不足时，系统应明确说不知道，而不是用推测补全。")

    add_section_title(doc, "3 目标用户与使用场景")
    add_subheading(doc, "3.1 核心用户")
    add_table(
        doc,
        ["用户类型", "典型任务", "主要困难"],
        [
            ["毕业论文写作者", "整理理论、方法和结论", "阅读周期长，前期材料难以回找"],
            ["课程论文写作者", "快速寻找论据和反例", "多篇论文之间缺少统一索引"],
            ["开题与综述写作者", "建立主题脉络和引用关系", "候选论文多，证据质量难判断"],
            ["跨语言研究者", "用中文问题检索英文论文", "术语和表述存在语言转换成本"],
        ],
        widths=[1.45, 2.25, 3.1],
    )

    add_subheading(doc, "3.2 核心场景")
    add_table(
        doc,
        ["场景", "用户动作", "系统结果"],
        [
            ["建立记忆", "导入 PDF、补充标签和读后笔记", "论文与笔记进入可检索文献库"],
            ["查找论据", "用自然语言提出写作问题", "返回相关论文、来源类型和命中片段"],
            ["单篇查证", "询问某篇论文是否包含指定信息", "明确回答有或没有，并展示原句与上下文"],
            ["扩展检索", "要求寻找文献库外论文", "查询 OpenAlex 和 Crossref，列出候选论文"],
            ["补全全文", "加入在线候选并导入开放 PDF", "全文可检索，失败时保留元数据和来源链接"],
        ],
        widths=[1.15, 2.5, 3.15],
    )

    add_section_title(doc, "4 快速开始")
    add_subheading(doc, "4.1 启动应用")
    add_paragraph(doc, "在项目目录双击 start.command，等待终端显示本地地址，浏览器会自动打开产品页面。运行期间不要关闭启动服务器所在的终端窗口。")
    add_bullets(doc, [
        "本地地址默认是 http://127.0.0.1:8765/。",
        "不要直接双击 index.html，PDF 解析模块需要通过本地服务器加载。",
        "论文数据默认保存在当前浏览器的 IndexedDB 中，不同浏览器之间不会自动同步。",
    ])

    add_subheading(doc, "4.2 导入第一篇论文")
    add_table(
        doc,
        ["步骤", "操作", "预期结果"],
        [
            ["1", "点击新对话左侧的导入论文", "打开 PDF、文本或 JSON 导入窗口"],
            ["2", "选择 PDF 并等待解析", "系统提取全文，显示字符数"],
            ["3", "补充作者、年份、主题标签", "论文信息进入文献库并参与检索"],
            ["4", "添加读后笔记和关键证据", "个人理解与论文全文建立关联"],
            ["5", "打开论文详情核对全文状态", "应显示全文已保存或待补 PDF"],
        ],
        widths=[0.55, 2.65, 3.6],
    )

    add_subheading(doc, "4.3 第一次提问")
    add_paragraph(doc, "在对话输入框中提出具体的写作问题，例如哪些论文讨论了二语经验对语音感知的影响。系统会先检索本地论文、笔记、标注和全文；本地证据不足时，智能联网模式会继续查询公开学术数据库。")

    add_section_title(doc, "5 核心功能说明")
    add_subheading(doc, "5.1 论文导入与全文提取")
    add_paragraph(doc, "产品支持 PDF、TXT、Markdown 和 JSON。PDF 导入通过 PDF.js 提取可复制文字，保存后全文会参与检索。扫描版 PDF 或图片型 PDF 没有可复制文字，当前版本不能自动识别，需要先进行 OCR。")
    add_paragraph(doc, "导入成功后，论文详情会显示全文状态。详情页只预览前 5,000 字以控制页面性能，但完整文本已经保存并用于全文查证。")

    add_subheading(doc, "5.2 读后笔记与证据片段")
    add_paragraph(doc, "读后笔记用于记录自己的理解、评价和可使用的写作场景。关键证据片段用于保存原文句子、研究方向或可支持的具体观点。系统在回答时会区分读后笔记、阅读标注、原文片段和自动提取全文。")

    add_subheading(doc, "5.3 文献库检索")
    add_paragraph(doc, "文献库支持按标题、作者、主题标签、摘要、笔记、证据片段和全文进行加权检索。标题和标签权重较高，个人笔记和证据片段优先于普通全文匹配，以减少无关段落干扰。")

    add_subheading(doc, "5.4 对话问答")
    add_paragraph(doc, "用户提出自然语言问题后，系统先给出本地证据，再根据问题意图判断是否进行在线扩展。回答卡片会显示论文标题、作者、年份、来源类型和命中片段，用户可以打开论文详情继续核对或补充笔记。")

    add_subheading(doc, "5.5 单篇论文全文查证")
    add_paragraph(doc, "在论文详情中，用户可以在“在这篇论文里查证”输入具体问题。系统执行中英文术语扩展，在指定论文全文内查找相关句子，并返回命中次数、原句和上下文。若论文没有全文，系统会明确提示无法判断正文。")
    add_table(
        doc,
        ["用户问题", "系统行为", "结果呈现"],
        [
            ["这篇论文里有没有提到粤语和普通话的重音差异", "扩展为 Cantonese、Mandarin、stress、perception 等词", "回答有，并列出命中片段"],
            ["论文里有没有提到某个变量", "在标题、摘要、正文和方法段落中检索", "显示原句、上下文和命中关键词"],
            ["关键词没有任何命中", "不直接判断论文绝对没有讨论", "说明关键词未命中并建议换词"],
            ["论文只有摘要", "只允许基于摘要回答", "明确说明不能确认正文"],
        ],
        widths=[2.2, 2.2, 2.4],
    )

    add_subheading(doc, "5.6 在线扩展检索")
    add_paragraph(doc, "当本地证据不足，或者用户明确提出寻找论文、跨语言研究、其他语言等需求时，系统会生成适合学术数据库的检索式，并同时查询 OpenAlex 与 Crossref。候选结果按标题、摘要、出版信息和被引数合并、去重和排序。")
    add_paragraph(doc, "用户可以把候选论文加入文献库，获得标题、作者、年份、摘要、DOI 和来源链接。如果候选提供开放 PDF，可以尝试自动下载和提取全文；若跨站限制导致失败，元数据仍会保存，并标记为待补 PDF。")

    add_subheading(doc, "5.7 数据导出与备份")
    add_paragraph(doc, "用户可以导出全部论文、笔记和证据为 JSON 文件，再在其他浏览器或未来版本中重新导入。导出文件包含个人资料，应妥善保管，不建议公开分享原始备份。")

    add_section_title(doc, "6 证据可信度规则")
    add_paragraph(doc, "产品将证据分为不同层级。回答必须尽量使用最高可信层级，并明确说明当前层级的限制。摘要不能替代全文，关键词未命中也不能证明论文完全没有讨论。")
    add_table(
        doc,
        ["证据层级", "适用情况", "回答约束"],
        [
            ["PDF 全文", "已经成功提取并保存正文", "可进行全文查证和上下文检索"],
            ["用户读后笔记", "记录个人理解和写作用途", "必须标注为个人笔记而非原文"],
            ["标注证据", "保存指定句子、页码或主题", "优先用于支持具体观点"],
            ["在线摘要", "只有公开摘要，没有 PDF", "只能基于摘要回答，并说明局限"],
            ["元数据", "只有标题、作者、年份和 DOI", "只能用于发现论文，不能判断结论"],
            ["待补 PDF", "已加入论文但未获得全文", "提示用户下载 PDF 后手动导入"],
        ],
        widths=[1.15, 2.55, 3.1],
    )

    add_section_title(doc, "7 数据与隐私")
    add_bullets(doc, [
        "论文、笔记、证据和对话默认保存在浏览器本地 IndexedDB 中。",
        "本地资料不会自动上传到产品服务器；在线检索时只发送检索关键词到 OpenAlex 和 Crossref。",
        "开放 PDF 自动导入会访问论文来源地址，是否允许跨站读取取决于来源服务器和浏览器策略。",
        "清除浏览器网站数据会同时删除本地文献库，因此需要定期导出 JSON 备份。",
        "导入他人论文 PDF 时需要遵守版权和数据库使用规则，产品不提供绕过付费墙的能力。",
    ])

    add_section_title(doc, "8 异常处理与常见问题")
    add_table(
        doc,
        ["问题", "原因", "处理方法"],
        [
            ["PDF 导入失败并提示 worker 服务器离线", "启动服务器后关闭了终端，或本地服务已停止", "双击 start.command，刷新页面后重新导入"],
            ["PDF 没有提取到文字", "扫描版或图片型 PDF", "先进行 OCR，或把关键段落手动录入证据片段"],
            ["在线检索返回 429", "公开学术接口限流", "稍后重试，或使用结果区提供的外部搜索入口"],
            ["开放 PDF 无法导入", "来源服务器不允许跨站读取", "打开来源下载 PDF，再在本应用中手动导入"],
            ["回答说没有找到", "本地没有匹配论文或关键词表述不同", "切换到智能联网，或改用标题、英文术语和变量名"],
            ["全文预览只显示前 5,000 字", "预览限制，不是全文缺失", "查看字符数并使用全文查证输入具体问题"],
        ],
        widths=[1.75, 2.3, 2.75],
    )

    add_section_title(doc, "9 产品边界")
    add_paragraph(doc, "当前版本使用关键词、中英文术语扩展、字段权重和上下文片段进行检索，尚未接入向量语义检索。完全不同措辞可能无法召回，因此系统应通过明确提示和换词建议管理用户预期。")
    add_bullets(doc, [
        "不替代 Zotero 等专业文献管理工具。",
        "不自动生成完整论文，不伪造引用，不把摘要冒充全文。",
        "不处理扫描版 PDF 的 OCR。",
        "不提供多人协作、云端同步和移动端原生应用。",
        "在线检索依赖公开学术接口，结果覆盖范围受数据源限制。",
    ])

    add_section_title(doc, "10 验收清单")
    add_table(
        doc,
        ["验收项", "通过标准"],
        [
            ["论文导入", "至少导入一篇 PDF，刷新后仍能打开并检索"],
            ["笔记保存", "新增读后笔记后可以在文献库和问答中检索"],
            ["本地问答", "问题至少返回一条可追溯依据"],
            ["单篇查证", "输入具体问题后能得到有或没有的明确结论"],
            ["在线检索", "候选论文显示题目、作者、年份、DOI 和 PDF 状态"],
            ["候选入库", "加入元数据后显示在线来源和待补 PDF状态"],
            ["异常恢复", "服务器断开、接口限流或 PDF 失败时给出明确处理路径"],
            ["数据导出", "导出的 JSON 可以重新导入并恢复资料"],
        ],
        widths=[1.45, 5.35],
    )

    add_section_title(doc, "11 版本与路线")
    add_table(
        doc,
        ["版本", "重点", "状态"],
        [
            ["v0.1", "本地论文、笔记和全文记忆", "已实现"],
            ["v0.2", "可追溯回答和单篇全文查证", "已实现"],
            ["v0.3", "OpenAlex、Crossref、开放 PDF 和待补 PDF", "已实现"],
            ["v0.4", "证据边界、错误恢复和中英文术语扩展", "已实现"],
            ["v1.0", "公开 Demo、项目案例页、测试数据和说明书", "当前版本"],
            ["v1.1", "向量检索、Zotero、BibTeX 和划词批注", "规划中"],
        ],
        widths=[0.85, 4.65, 1.3],
    )

    add_section_title(doc, "12 文档维护信息")
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["文档名称", "论文记忆与引用助手产品说明书"],
            ["文档版本", "v1.0"],
            ["最后更新", "2026年9月"],
            ["维护方式", "功能范围或证据规则发生变化时同步更新"],
            ["配套材料", "PRD、用户流程、验证方案、Demo 脚本和项目案例页"],
        ],
        widths=[1.3, 5.5],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
